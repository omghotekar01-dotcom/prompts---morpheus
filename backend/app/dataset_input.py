from __future__ import annotations

"""Spreadsheet/CSV I/O adapters. Native C++ remains responsible for profiling."""

import base64
import csv
import hashlib
import io
import keyword
import re
from dataclasses import dataclass


MAX_CSV_BYTES = 15_000_000
MAX_SPREADSHEET_BYTES = 15_000_000
# A generated field becomes a C++ struct member, not just a Python identifier.
_CPP_RESERVED = frozenset(
    "alignas alignof and and_eq asm auto bitand bitor bool break case catch "
    "char char8_t char16_t char32_t class compl concept const consteval constexpr "
    "constinit const_cast continue co_await co_return co_yield decltype default "
    "delete do double dynamic_cast else enum explicit export extern false float "
    "for friend goto if inline int long mutable namespace new noexcept not not_eq "
    "nullptr operator or or_eq private protected public register reinterpret_cast "
    "requires return short signed sizeof static static_assert static_cast struct "
    "switch template this thread_local throw true try typedef typeid typename "
    "union unsigned using virtual void volatile wchar_t while xor xor_eq"
    .split()
)


@dataclass(frozen=True)
class PreparedDataset:
    csv_text: str
    source_columns: list[str]
    canonical_columns: list[str]
    source_sha256: str
    source_kind: str
    worksheet: str | None = None


def normalize_csv(source: str, *, source_kind: str = "CSV", worksheet: str | None = None) -> PreparedDataset:
    if not isinstance(source, str) or not source.strip():
        raise ValueError("Supply a CSV header and at least one record")
    if len(source.encode("utf-8")) > MAX_CSV_BYTES:
        raise ValueError("CSV exceeds the 15 MB import limit")
    reader = csv.reader(io.StringIO(source.lstrip("\ufeff"), newline=""), strict=True)
    try:
        header = next(reader)
    except (StopIteration, csv.Error) as exc:
        raise ValueError("Cannot read CSV header") from exc
    if not 1 <= len(header) <= 64:
        raise ValueError("Dataset requires 1-64 columns")
    names: list[str] = []
    used: set[str] = set()
    for index, raw in enumerate(header, 1):
        normalized = re.sub(r"[^A-Za-z0-9_]", "_", raw.strip())
        normalized = re.sub(r"_+", "_", normalized).strip("_")
        if not normalized or not re.match(r"^[A-Za-z_]", normalized):
            normalized = "field_" + normalized if normalized else f"field_{index}"
        if (normalized in _CPP_RESERVED or keyword.iskeyword(normalized)
                or normalized.startswith("__") or len(normalized) > 100):
            normalized = f"field_{index}"
        proposed = normalized
        suffix = 2
        while proposed in used:
            proposed = f"{normalized}_{suffix}"
            suffix += 1
        used.add(proposed)
        names.append(proposed)

    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(names)
    count = 0
    try:
        for row in reader:
            if len(row) != len(names):
                raise ValueError(f"CSV record {count + 1} has {len(row)} columns; expected {len(names)}")
            writer.writerow(row)
            count += 1
            if count > 100_000:
                raise ValueError("CSV exceeds the 100,000-record cap")
    except csv.Error as exc:
        raise ValueError("CSV formatting error: " + str(exc)) from exc
    if count == 0:
        raise ValueError("CSV needs at least one data record")
    processed = output.getvalue()
    if len(processed.encode("utf-8")) > MAX_CSV_BYTES:
        raise ValueError("Normalized CSV exceeds the 15 MB limit")
    return PreparedDataset(
        csv_text=processed, source_columns=header, canonical_columns=names,
        source_sha256=hashlib.sha256(source.encode("utf-8")).hexdigest(),
        source_kind=source_kind, worksheet=worksheet,
    )


def excel_to_csv(encoded: str) -> PreparedDataset:
    try:
        binary = base64.b64decode(encoded, validate=True)
    except (ValueError, base64.binascii.Error) as exc:
        raise ValueError("Invalid base64 Excel workbook") from exc
    if len(binary) > MAX_SPREADSHEET_BYTES:
        raise ValueError("Excel workbook exceeds the 15 MB import limit")
    if not binary.startswith(b"PK"):
        raise ValueError("Only .xlsx Excel workbooks are supported; legacy .xls is not supported")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel support requires openpyxl; install backend requirements") from exc
    try:
        workbook = load_workbook(io.BytesIO(binary), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Unable to read .xlsx workbook") from exc
    try:
        sheet = workbook.active
        worksheet = sheet.title
        output = io.StringIO(newline="")
        writer = csv.writer(output, lineterminator="\n")
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index > 100_000:
                raise ValueError("Worksheet exceeds the 100,000-record cap")
            if len(row) > 64:
                raise ValueError("Worksheet exceeds the 64-column cap")
            values = ["" if value is None else str(value) for value in row]
            writer.writerow(values)
            if output.tell() > MAX_CSV_BYTES:
                raise ValueError("Worksheet CSV exceeds the 15 MB limit")
        prepared = normalize_csv(output.getvalue(), source_kind="XLSX", worksheet=worksheet)
        return PreparedDataset(
            csv_text=prepared.csv_text,
            source_columns=prepared.source_columns,
            canonical_columns=prepared.canonical_columns,
            source_sha256=hashlib.sha256(binary).hexdigest(),
            source_kind="XLSX",
            worksheet=worksheet,
        )
    finally:
        workbook.close()
